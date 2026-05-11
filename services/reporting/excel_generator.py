"""Excel report generator using openpyxl.

Generates multi-sheet Excel reports with cleaned data, statistics,
anomalies, and AI insights.
"""

import logging
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

from config.settings import EXPORT_DIR
from repositories.dataset_repository import DatasetRepository

logger = logging.getLogger(__name__)

_dataset_repo = DatasetRepository()


def generer(
    dataset_id: int,
    user_id: int,
    analytics: dict,
    anomalies: dict,
    insights: str,
) -> str:
    """Generate an Excel report for a dataset analysis.

    Creates a multi-sheet Excel workbook with:
    - Sheet 1: Données nettoyées (cleaned data)
    - Sheet 2: Statistiques descriptives
    - Sheet 3: Anomalies détectées
    - Sheet 4: Résumé IA (insights text)

    Args:
        dataset_id: The dataset ID.
        user_id: The user who requested the report.
        analytics: Analytics results dictionary.
        anomalies: Anomalies detection results dictionary.
        insights: LLM-generated insights text.

    Returns:
        The file path to the generated Excel file.
    """
    # Prepare output path
    export_dir = Path(EXPORT_DIR) / "reports"
    export_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"rapport_dataset_{dataset_id}_{timestamp}.xlsx"
    output_path = export_dir / filename

    logger.info("Génération du rapport Excel: %s", output_path)

    # Load dataset for cleaned data
    dataset = _dataset_repo.find_by_id(dataset_id)

    wb = Workbook()

    # Define styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # --- Sheet 1: Données nettoyées ---
    ws_data = wb.active
    ws_data.title = "Données nettoyées"

    if dataset and dataset.chemin_fichier:
        try:
            df = pd.read_csv(dataset.chemin_fichier, encoding="utf-8")
            # Limit to first 10000 rows for Excel performance
            df_export = df.head(10000)

            for r_idx, row in enumerate(dataframe_to_rows(df_export, index=False, header=True), 1):
                for c_idx, value in enumerate(row, 1):
                    cell = ws_data.cell(row=r_idx, column=c_idx, value=value)
                    if r_idx == 1:
                        cell.font = header_font
                        cell.fill = header_fill
                        cell.alignment = header_alignment
                    cell.border = thin_border

            # Auto-adjust column widths
            for col in ws_data.columns:
                max_length = 0
                for cell in col[:50]:  # Check first 50 rows
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                adjusted_width = min(max_length + 2, 30)
                ws_data.column_dimensions[col[0].column_letter].width = adjusted_width

            if len(df) > 10000:
                ws_data.cell(
                    row=len(df_export) + 2,
                    column=1,
                    value=f"[Affichage limité aux 10 000 premières lignes sur {len(df)} total]",
                )
        except Exception as e:
            logger.warning("Erreur lecture données pour Excel: %s", e)
            ws_data.cell(row=1, column=1, value="Données non disponibles")
    else:
        ws_data.cell(row=1, column=1, value="Dataset non trouvé")

    # --- Sheet 2: Statistiques descriptives ---
    ws_stats = wb.create_sheet("Statistiques descriptives")
    stats_data = analytics.get("statistiques", {}).get("colonnes", {})

    if stats_data:
        # Header row
        headers = ["Colonne", "Moyenne", "Médiane", "Écart-type", "Min", "Max", "Q25", "Q75", "Skewness", "Kurtosis", "Count"]
        for c_idx, header in enumerate(headers, 1):
            cell = ws_stats.cell(row=1, column=c_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Data rows
        for r_idx, (col_name, col_stats) in enumerate(stats_data.items(), 2):
            values = [
                col_name,
                col_stats.get("mean"),
                col_stats.get("median"),
                col_stats.get("std"),
                col_stats.get("min"),
                col_stats.get("max"),
                col_stats.get("q25"),
                col_stats.get("q75"),
                col_stats.get("skewness"),
                col_stats.get("kurtosis"),
                col_stats.get("count"),
            ]
            for c_idx, value in enumerate(values, 1):
                cell = ws_stats.cell(row=r_idx, column=c_idx, value=value)
                cell.border = thin_border
                if c_idx > 1 and value is not None:
                    cell.number_format = "0.0000"
    else:
        ws_stats.cell(row=1, column=1, value="Aucune statistique disponible")

    # --- Sheet 3: Anomalies détectées ---
    ws_anomalies = wb.create_sheet("Anomalies détectées")

    all_anomalies = (
        anomalies.get("zscore", [])
        + anomalies.get("iqr", [])
        + anomalies.get("isolation", [])
    )
    all_anomalies.sort(key=lambda x: x.get("score", 0), reverse=True)

    # Header
    anom_headers = ["Ligne", "Colonne", "Score", "Algorithme", "Valeur"]
    for c_idx, header in enumerate(anom_headers, 1):
        cell = ws_anomalies.cell(row=1, column=c_idx, value=header)
        cell.font = header_font
        cell.fill = PatternFill(start_color="C62828", end_color="C62828", fill_type="solid")
        cell.alignment = header_alignment
        cell.border = thin_border

    # Data
    for r_idx, anomaly in enumerate(all_anomalies[:5000], 2):
        values = [
            anomaly.get("ligne"),
            anomaly.get("colonne"),
            anomaly.get("score"),
            anomaly.get("type"),
            anomaly.get("valeur"),
        ]
        for c_idx, value in enumerate(values, 1):
            cell = ws_anomalies.cell(row=r_idx, column=c_idx, value=value)
            cell.border = thin_border

    # Summary row
    summary_row = len(all_anomalies[:5000]) + 3
    ws_anomalies.cell(row=summary_row, column=1, value="Résumé").font = Font(bold=True)
    ws_anomalies.cell(row=summary_row + 1, column=1, value="Total anomalies:")
    ws_anomalies.cell(row=summary_row + 1, column=2, value=anomalies.get("total", 0))
    resume = anomalies.get("resume", {})
    ws_anomalies.cell(row=summary_row + 2, column=1, value="Z-Score:")
    ws_anomalies.cell(row=summary_row + 2, column=2, value=resume.get("zscore_count", 0))
    ws_anomalies.cell(row=summary_row + 3, column=1, value="IQR:")
    ws_anomalies.cell(row=summary_row + 3, column=2, value=resume.get("iqr_count", 0))
    ws_anomalies.cell(row=summary_row + 4, column=1, value="Isolation Forest:")
    ws_anomalies.cell(row=summary_row + 4, column=2, value=resume.get("isolation_count", 0))

    # --- Sheet 4: Résumé IA ---
    ws_insights = wb.create_sheet("Résumé IA")

    ws_insights.cell(row=1, column=1, value="Insights générés par Intelligence Artificielle")
    ws_insights.cell(row=1, column=1).font = Font(size=14, bold=True)

    ws_insights.cell(row=2, column=1, value=f"Date: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    ws_insights.cell(row=3, column=1, value="")

    if insights:
        for r_idx, line in enumerate(insights.split("\n"), 4):
            ws_insights.cell(row=r_idx, column=1, value=line.strip())
            ws_insights.column_dimensions["A"].width = 100
    else:
        ws_insights.cell(row=4, column=1, value="Aucun insight IA disponible.")

    # Save workbook
    try:
        wb.save(str(output_path))
        logger.info("Rapport Excel généré avec succès: %s", output_path)
    except Exception as e:
        logger.error("Erreur lors de la sauvegarde du rapport Excel: %s", e)
        raise

    return str(output_path)
